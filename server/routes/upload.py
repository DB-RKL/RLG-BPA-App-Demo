import asyncio
import io

from fastapi import APIRouter, UploadFile, File, HTTPException

from ..db import db
from ..models import DocumentOut

router = APIRouter(tags=["upload"])


def _pdf_sync(content: bytes) -> str:
    import pdfplumber

    text_parts: list[str] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)

            tables = page.extract_tables()
            for table in tables:
                for row in table:
                    cells = [str(c) if c else "" for c in row]
                    text_parts.append(" | ".join(cells))

    return "\n\n".join(text_parts)


def _excel_sync(content: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), data_only=True)
    text_parts: list[str] = []

    for sheet in wb.worksheets:
        text_parts.append(f"--- Sheet: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                text_parts.append(" | ".join(cells))

    return "\n".join(text_parts)


def _csv_sync(content: bytes) -> str:
    import csv

    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError:
        text = content.decode("latin-1")

    reader = csv.reader(io.StringIO(text))
    text_parts: list[str] = []
    for row in reader:
        cells = [c.strip() for c in row]
        if any(cells):
            text_parts.append(" | ".join(cells))

    return "\n".join(text_parts)


def _docx_sync(content: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(content))
    text_parts: list[str] = []

    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                text_parts.append(" | ".join(cells))

    return "\n\n".join(text_parts)


# CPU-bound parsers are run on the default thread pool via `asyncio.to_thread`
# so uvicorn's event loop stays responsive during batch extraction.
async def extract_text_from_pdf(content: bytes) -> str:
    return await asyncio.to_thread(_pdf_sync, content)


async def extract_text_from_excel(content: bytes) -> str:
    return await asyncio.to_thread(_excel_sync, content)


async def extract_text_from_csv(content: bytes) -> str:
    return await asyncio.to_thread(_csv_sync, content)


async def extract_text_from_docx(content: bytes) -> str:
    return await asyncio.to_thread(_docx_sync, content)


EXT_TO_TYPE = {
    "pdf": "pdf",
    "xlsx": "excel",
    "xls": "excel",
    "csv": "csv",
    "docx": "docx",
}

EXTRACTORS = {
    "pdf": extract_text_from_pdf,
    "excel": extract_text_from_excel,
    "csv": extract_text_from_csv,
    "docx": extract_text_from_docx,
}


@router.post("/upload", response_model=DocumentOut)
async def upload_document(file: UploadFile = File(...)):
    if not file.filename:
        raise HTTPException(400, "No filename provided")

    ext = file.filename.rsplit(".", 1)[-1].lower()
    file_type = EXT_TO_TYPE.get(ext)
    if not file_type:
        raise HTTPException(400, f"Unsupported file type: .{ext}. Use PDF, Excel, CSV, or Word.")

    content = await file.read()

    try:
        raw_text = await EXTRACTORS[file_type](content)
    except Exception as e:
        raise HTTPException(422, f"Failed to extract text: {e}")

    if not raw_text.strip():
        raise HTTPException(422, "No text could be extracted from the file")

    pool = await db.get_pool()
    if not pool:
        raise HTTPException(503, "Database not available")

    kind_expr = (
        "document_kind" if db.has_document_kind else "NULL::text AS document_kind"
    )
    async with pool.acquire() as conn:
        row = await conn.fetchrow(
            f"""
            INSERT INTO documents (filename, file_type, status, raw_text)
            VALUES ($1, $2, 'uploaded', $3)
            RETURNING id, filename, file_type, upload_time, status, raw_text,
                      {kind_expr}
            """,
            file.filename,
            file_type,
            raw_text,
        )

    return DocumentOut(**dict(row))
